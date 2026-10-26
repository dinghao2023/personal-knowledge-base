"""入库链路：LangChain 加载文档 → 分块 → 向量化 → 写入向量库。"""
import os
from datetime import datetime
from typing import Iterator, List

from langchain_core.document_loaders import BaseLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter

from .config import settings
from .db import get_collection, get_vectorstore

SUPPORTED_EXTS = {"pdf", "docx", "md", "txt"}

# 断点顺序与原先手写分块一致：先段落，再换行，再中英文句末
_SEPARATORS = ["\n\n", "\n", "。", "！", "？", ".", "!", "；", " ", ""]


class KnowledgeFileLoader(BaseLoader):
    """把一份已落盘文件载成 LangChain Document。PDF 按页保留从 1 开始的页码。"""

    def __init__(self, path: str, filename: str):
        self.path = path
        self.filename = filename

    def lazy_load(self) -> Iterator[Document]:
        ext = self.path.rsplit(".", 1)[-1].lower()
        meta = {"source": self.filename, "page": 0}
        if ext == "pdf":
            import fitz

            doc = fitz.open(self.path)
            try:
                for i, page in enumerate(doc):
                    text = page.get_text()
                    if text.strip():
                        yield Document(
                            page_content=text,
                            metadata={"source": self.filename, "page": i + 1},
                        )
            finally:
                doc.close()
            return
        if ext == "docx":
            import docx

            paras = [p.text for p in docx.Document(self.path).paragraphs if p.text]
            text = "\n".join(paras)
        elif ext in {"md", "txt"}:
            with open(self.path, encoding="utf-8") as f:
                text = f.read()
        else:
            raise ValueError(f"不支持的文件类型: .{ext}")
        if text.strip():
            yield Document(page_content=text, metadata=meta)


def _splitter() -> RecursiveCharacterTextSplitter:
    return RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        separators=_SEPARATORS,
    )


def _delete_stale_chunks(filename: str, keep_ids: set[str]) -> None:
    """删掉同一来源里这次没有写入的分块。新文件更短时，多出来的旧序号会留在这里。"""
    data = get_collection().get(where={"source": filename}, include=["metadatas"])
    stale = [item_id for item_id in data.get("ids") or [] if item_id not in keep_ids]
    if stale:
        get_collection().delete(ids=stale)


def ingest_file(file_path: str, filename: str) -> int:
    """解析文档并写入向量库，返回分块数量。

    同一文件再次上传时，相同序号会被覆盖。写入成功后再删掉本次没有覆盖到的旧分块。
    抽不出文字时清掉该文件的全部分块，避免继续检索到上一版内容。
    """
    docs = _splitter().split_documents(list(KnowledgeFileLoader(file_path, filename).lazy_load()))
    docs = [doc for doc in docs if doc.page_content.strip()]
    if not docs:
        _delete_stale_chunks(filename, set())
        return 0

    ids = [f"{filename}_{i}" for i in range(len(docs))]
    get_vectorstore().add_documents(docs, ids=ids)
    _delete_stale_chunks(filename, set(ids))
    return len(docs)


def _chunk_counts() -> dict:
    """按文件名统计向量库中的分块数量。"""
    counts: dict = {}
    try:
        data = get_collection().get(include=["metadatas"])
    except Exception:
        return counts
    for meta in data.get("metadatas") or []:
        source = (meta or {}).get("source")
        if source:
            counts[source] = counts.get(source, 0) + 1
    return counts


def list_saved_files() -> List[dict]:
    """列出 docs 目录中已保存的文件，并附上分块数与修改时间。"""
    os.makedirs(settings.docs_dir, exist_ok=True)
    counts = _chunk_counts()
    files: List[dict] = []
    for name in os.listdir(settings.docs_dir):
        if name.startswith("."):
            continue
        path = os.path.join(settings.docs_dir, name)
        if not os.path.isfile(path):
            continue
        stat = os.stat(path)
        uploaded_at = datetime.fromtimestamp(stat.st_mtime).astimezone().isoformat(timespec="seconds")
        files.append(
            {
                "filename": name,
                "size": stat.st_size,
                "chunks": counts.get(name, 0),
                "uploaded_at": uploaded_at,
            }
        )
    files.sort(key=lambda item: item["uploaded_at"], reverse=True)
    return files


def delete_saved_file(filename: str) -> None:
    """删除 docs 目录中的文件，并清除向量库里该文件的全部分块。"""
    name = os.path.basename(filename)
    if not name or name != filename or name.startswith("."):
        raise ValueError("非法文件名")

    os.makedirs(settings.docs_dir, exist_ok=True)
    docs_root = os.path.realpath(settings.docs_dir)
    path = os.path.realpath(os.path.join(docs_root, name))
    if os.path.commonpath([docs_root, path]) != docs_root:
        raise ValueError("非法文件名")
    if not os.path.isfile(path):
        raise FileNotFoundError(name)

    get_collection().delete(where={"source": name})
    os.remove(path)


def save_upload(upload_file, filename: str) -> str:
    """保存上传文件到 docs 目录，返回落盘路径。"""
    os.makedirs(settings.docs_dir, exist_ok=True)
    dest = os.path.join(settings.docs_dir, filename)
    with open(dest, "wb") as f:
        f.write(upload_file.file.read())
    return dest
